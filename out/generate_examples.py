# ruff: noqa: RUF001
"""Create realistic procurement inputs. No customer data or live answers are generated.

The product photo is an unmodified image from ekt.kz, with provenance in manifest.
All business documents are constructed examples, using actual catalogue references.
"""

import csv
import hashlib
import io
import json
import os
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import pymupdf
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent
INPUTS = OUT / "inputs"
FONT = Path(os.getenv("EXAMPLE_FONT", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
PRODUCTS = {
    p["id"]: p
    for f in (ROOT / "docs/source of truth/assets").glob("products*.json")
    for p in json.loads(f.read_text())["items"]
}


def item(pid, quantity):
    p = PRODUCTS[pid]
    return {
        "id": pid,
        "article": p["article"],
        "code": p["name"].split()[0] if pid != 45357 else p["article"],
        "name": p["name"],
        "quantity": quantity,
    }


CASES = [
    {
        "path": "inputs/purchase_request.xlsx",
        "scenario": "Заявка снабжения из ERP: два листа, внутренние коды, повтор товара, вычисленное количество",
        "uc": ["UC-01", "UC-04", "UC-05"],
        "prompt": "Это заявка на закупку. Проверь все позиции и количества на обоих листах по вашему каталогу.",
        "expected_items": [
            item(515280, 2),
            item(515281, 1),
            item(515280, 1),
            item(515285, 2),
        ],
        "confirm_cart": True,
    },
    {
        "path": "inputs/procurement_request.docx",
        "scenario": "Письмо отдела снабжения: позиции в тексте и таблице, количество одной позиции уточняется",
        "uc": ["UC-01", "UC-04"],
        "prompt": "Проверь товары из письма: наличие, характеристики и что нужно уточнить перед заказом.",
        "expected_items": [item(515279, 3), item(515283, 2), item(515291, None)],
    },
    {
        "path": "inputs/project_specification.pdf",
        "scenario": "Двухстраничная проектная спецификация: позиции QF/EL и отсутствующий товар",
        "uc": ["UC-01", "UC-02"],
        "prompt": "Проверь спецификацию с обеих страниц. Покажи наличие всех позиций и аналоги отсутствующих.",
        "expected_items": [item(515280, 2), item(515281, 1), item(45357, 4)],
    },
    {
        "path": "inputs/delivery_note_scan.pdf",
        "scenario": "Скан накладной для повторной закупки, без текстового слоя",
        "uc": ["UC-01", "UC-04"],
        "prompt": "Нужно докупить те же позиции в тех же количествах, что в этой накладной. Проверь по каталогу.",
        "expected_items": [item(515283, 2), item(515285, 1)],
    },
    {
        "path": "inputs/delivery_note_image.jpg",
        "scenario": "JPEG страницы накладной из мессенджера: две позиции, без текстового слоя",
        "uc": ["UC-01", "UC-04"],
        "prompt": "Это изображение накладной. Найди обе позиции, нужно докупить столько же штук.",
        "expected_items": [item(515283, 2), item(515285, 1)],
        "origin": "Raster rendering of the constructed delivery note; not a customer photograph",
    },
    {
        "path": "inputs/product_photo.jpg",
        "scenario": "Фото автомата с мелкой неоднозначной маркировкой: уточнение точного кода покупателем",
        "uc": ["UC-01", "UC-04"],
        "prompt": "Найди такой же автомат по маркировке на фото. Нужны 2 штуки. Покажи характеристики и наличие.",
        "expected_items": [item(515280, 2)],
        "origin": "Unmodified ekt.kz catalogue photograph, not a customer field photo",
        "source_url": PRODUCTS[515280]["image"],
        "marking_limitation": "The source URL identifies 027005, but the small photographed reference is not reliably readable. The URL is provenance, not visual ground truth.",
        "requires_clarification": True,
        "clarification": "Уточнил маркировку на самом аппарате: код производителя 027005. Нужны 2 штуки. Покажи этот товар и наличие, пока не добавляй.",
    },
    {
        "path": "inputs/erp_export.csv",
        "scenario": "CSV экспорта ERP: разделитель ;, внутренний код и цена стороннего предложения",
        "uc": ["UC-01", "UC-04"],
        "prompt": "Найди позиции этой выгрузки у вас. Цена в файле от другого поставщика; нужна ваша цена и наличие.",
        "expected_items": [item(515278, 4), item(515282, 2)],
        "external_prices": [1, 2],
    },
    {
        "path": "inputs/electrician_request.txt",
        "scenario": "Неформальная заявка электрика: числа словами, количество одной позиции не определено",
        "uc": ["UC-01", "UC-04"],
        "prompt": "Проверь заявку электрика. Найди все позиции и покажи, что ещё надо уточнить.",
        "expected_items": [item(515280, 2), item(515281, 1), item(515285, None)],
    },
]


def write_excel():
    sheets = [
        (
            "Основная заявка",
            [
                [
                    "Заявка снабжения № ЗС-0923",
                    "Объект: ЩР-1 / административный корпус",
                ],
                [
                    "Внутренний код ERP",
                    "Поз. проекта",
                    "Наименование",
                    "Код производителя",
                    "Ед.",
                    "Количество",
                    "Примечание",
                ],
                [
                    "00001234",
                    "QF1-QF2",
                    "Legrand DRX125 MT 3P 50A 10kA",
                    "027005",
                    "шт",
                    2,
                    "Основной щит",
                ],
                [
                    "00001235",
                    "QF3",
                    "Legrand DRX125 MT 3P 100A 20kA",
                    "027028",
                    "шт",
                    ("1*1", 1),
                    "Ввод",
                ],
                [],
                [
                    "Итого строк",
                    "",
                    "",
                    "",
                    "",
                    2,
                    "Остатки и цены запрашиваются у поставщика",
                ],
            ],
        ),
        (
            "Дополнительные материалы",
            [
                ["Дополнение к заявке ЗС-0923"],
                [
                    "Внутренний код ERP",
                    "Поз. проекта",
                    "Наименование",
                    "Код производителя",
                    "Ед.",
                    "Количество",
                    "Примечание",
                ],
                [
                    "00001234",
                    "QF4",
                    "Legrand DRX125 MT 3P 50A 10kA",
                    "027005",
                    "шт",
                    1,
                    "Повтор для второго щита, учесть дополнительно",
                ],
                [
                    "00001236",
                    "QF5-QF6",
                    "Legrand DRX125 MT 3P 63A 20kA",
                    "027220",
                    "шт",
                    2,
                    "Резервные линии",
                ],
                [
                    "Комментарий",
                    "Предложить доступные варианты, замену согласовать отдельно",
                ],
            ],
        ),
    ]
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    parts = {
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{ns}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
        + "".join(
            f'<sheet name="{name}" sheetId="{i}" r:id="rId{i}"/>'
            for i, (name, _) in enumerate(sheets, 1)
        )
        + "</sheets></workbook>",
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(
            f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>'
            for i in range(1, 3)
        )
        + "</Relationships>",
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        + "".join(
            f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            for i in range(1, 3)
        )
        + "</Types>",
    }
    for i, (_, rows) in enumerate(sheets, 1):
        rendered = []
        for n, row in enumerate(rows, 1):
            cells = []
            for col, value in zip("ABCDEFG", row, strict=False):
                ref = f"{col}{n}"
                if isinstance(value, tuple):
                    cells.append(f'<c r="{ref}"><f>{value[0]}</f><v>{value[1]}</v></c>')
                elif isinstance(value, int):
                    cells.append(f'<c r="{ref}"><v>{value}</v></c>')
                else:
                    cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(value)}</t></is></c>')
            rendered.append(f'<row r="{n}">{"".join(cells)}</row>')
        parts[f"xl/worksheets/sheet{i}.xml"] = (
            f'<worksheet xmlns="{ns}"><sheetViews><sheetView workbookViewId="0"><pane ySplit="2" topLeftCell="A3" state="frozen"/></sheetView></sheetViews><cols><col min="1" max="2" width="24" customWidth="1"/><col min="3" max="3" width="45" customWidth="1"/><col min="4" max="7" width="25" customWidth="1"/></cols><sheetData>'
            + "".join(rendered)
            + "</sheetData></worksheet>"
        )
    with zipfile.ZipFile(INPUTS / "purchase_request.xlsx", "w", zipfile.ZIP_DEFLATED) as z:
        for name, body in parts.items():
            z.writestr(name, '<?xml version="1.0" encoding="UTF-8"?>' + body)


def write_word():
    d = Document()
    d.core_properties.title = "Заявка отдела снабжения ЗС-0941"
    d.core_properties.author = "Отдел снабжения"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading("ЗАПРОС НАЛИЧИЯ И ХАРАКТЕРИСТИК", 0)
    d.add_paragraph(
        "Исх. ЗС-0941 от 23.09.2026\nОбъект: административный корпус, распределительные щиты"
    )
    d.add_paragraph(
        "Просим проверить следующие позиции Legrand. Указанные количества нужны для комплектации щитов; замену просим согласовать отдельно."
    )
    d.add_paragraph(
        "1. Автоматический выключатель DRX125 MT, 3P, 40 А, 10 kA, код производителя 027004 — 3 шт."
    )
    d.add_paragraph(
        "2. Автоматический выключатель DRX125 MT, 3P, 40 А, 20 kA, код производителя 027024 — 2 шт."
    )
    d.add_paragraph("3. Дополнительно требуется позиция:")
    table = d.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, value in zip(
        table.rows[0].cells,
        ["Код производителя", "Наименование", "Количество"],
        strict=True,
    ):
        cell.text = value
    for cell, value in zip(
        table.add_row().cells,
        ["027228", "DRX250 MT, 3P, 160 А, 18 kA", "Уточняется"],
        strict=True,
    ):
        cell.text = value
    d.add_paragraph(
        "Потребуется документация на изделия. Просим отдельно сообщить условия оплаты и доставки. Желаемая дата получения — до конца недели; это пожелание заказчика, а не согласованный срок."
    )
    d.add_paragraph("Отдел снабжения")
    for section in d.sections:
        section.left_margin = section.right_margin = Inches(0.7)
    d.save(INPUTS / "procurement_request.docx")


def page_base(pdf, title, reference, page_number):
    page = pdf.new_page(width=842, height=595)
    page.insert_font(fontname="body", fontfile=str(FONT))
    page.insert_text((36, 42), title, fontname="body", fontsize=18)
    page.insert_text((36, 66), reference, fontname="body", fontsize=10)
    page.draw_line((36, 80), (806, 80), color=(0.25, 0.3, 0.35))
    page.insert_text(
        (36, 560),
        f"Отдел снабжения | 23.09.2026 | Лист {page_number}",
        fontname="body",
        fontsize=9,
    )
    return page


def table(page, rows):
    xs = [36, 108, 215, 670, 714, 806]
    entries = [["Поз.", "Код / артикул", "Наименование", "Ед.", "Кол-во"], *rows]
    for n, values in enumerate(entries):
        top = 108 + 70 * n
        for left, right, value in zip(xs[:-1], xs[1:], values, strict=True):
            page.draw_rect(
                pymupdf.Rect(left, top, right, top + 70),
                color=(0.55, 0.55, 0.55),
                fill=(0.92, 0.94, 0.96) if n == 0 else None,
            )
            result = page.insert_textbox(
                pymupdf.Rect(left + 6, top + 8, right - 6, top + 65),
                str(value),
                fontname="body",
                fontsize=10,
            )
            if result < 0:
                raise ValueError(f"PDF cell does not fit: {value}")


def write_pdfs():
    pdf = pymupdf.open()
    p = page_base(
        pdf,
        "СПЕЦИФИКАЦИЯ ЭЛЕКТРООБОРУДОВАНИЯ",
        "СП-0923 | Административный корпус | Раздел ЭОМ, щит ЩР-1",
        1,
    )
    table(
        p,
        [
            ["QF1-QF2", "027005", "Автомат Legrand DRX125 MT 3P 50A 10kA", "шт", 2],
            ["QF3", "027028", "Автомат Legrand DRX125 MT 3P 100A 20kA", "шт", 1],
        ],
    )
    p.insert_text(
        (36, 400),
        "Продолжение перечня на листе 2. QF и EL — проектные обозначения, не артикулы.",
        fontname="body",
        fontsize=10,
    )
    p = page_base(pdf, "СПЕЦИФИКАЦИЯ — ПРОДОЛЖЕНИЕ", "СП-0923 | Освещение общего помещения", 2)
    table(
        p,
        [["EL1-EL4", "ярп4520", "LED STARK 30W 2400Lm 4000K IP20 MEGALIGHT", "шт", 4]],
    )
    p.insert_textbox(
        pymupdf.Rect(36, 310, 800, 465),
        "При отсутствии предложить доступный аналог с указанием отличий. Подмена позиции без согласования не допускается. Цены и складские остатки в проектную спецификацию не входят.",
        fontname="body",
        fontsize=11,
    )
    pdf.save(INPUTS / "project_specification.pdf", deflate=True)
    pdf.close()
    paper = pymupdf.open()
    p = page_base(
        paper,
        "НАКЛАДНАЯ НА ВНУТРЕННЕЕ ПЕРЕМЕЩЕНИЕ",
        "НК-45821 | Со склада снабжения на участок сборки ЩР-2",
        1,
    )
    table(
        p,
        [
            ["1", "027024", "Автомат Legrand DRX125 MT 3P 40A 20kA", "шт", 2],
            ["2", "027220", "Автомат Legrand DRX125 MT 3P 63A 20kA", "шт", 1],
        ],
    )
    p.insert_text(
        (36, 415),
        "Отпустил: ____________           Получил: ____________",
        fontname="body",
        fontsize=11,
    )
    scan = pymupdf.open()
    raster = p.get_pixmap(matrix=pymupdf.Matrix(1.8, 1.8), alpha=False)
    (INPUTS / "delivery_note_image.jpg").write_bytes(raster.tobytes("jpeg", jpg_quality=85))
    scan.new_page(width=842, height=595).insert_image(
        p.rect, stream=raster.tobytes("jpeg", jpg_quality=85)
    )
    scan.save(INPUTS / "delivery_note_scan.pdf", deflate=True)
    scan.close()
    paper.close()


def main():
    INPUTS.mkdir(exist_ok=True)
    if not FONT.exists():
        raise SystemExit("Set EXAMPLE_FONT to a Cyrillic TTF font")
    if not (INPUTS / "product_photo.jpg").exists():
        raise SystemExit(
            "Download the unmodified photograph from the source_url in generate_examples.py before generating"
        )
    write_excel()
    write_word()
    write_pdfs()
    stream = io.StringIO(newline="")
    w = csv.writer(stream, delimiter=";")
    w.writerow(
        [
            "request_id",
            "internal_sku",
            "manufacturer_code",
            "product_name",
            "qty",
            "unit",
            "other_supplier_price",
            "comment",
        ]
    )
    w.writerow(
        [
            "REQ-8421",
            "00004521",
            "027008",
            "Legrand DRX125 MT 3P 100A 10kA",
            4,
            "шт",
            1,
            "Цена другого поставщика; перепроверить",
        ]
    )
    w.writerow(
        [
            "REQ-8421",
            "00004522",
            "027022",
            "Legrand DRX125 MT 3P 25A 20kA",
            2,
            "шт",
            2,
            "Щит; запасная линия",
        ]
    )
    (INPUTS / "erp_export.csv").write_text(stream.getvalue(), encoding="utf-8-sig")
    (INPUTS / "electrician_request.txt").write_text(
        "Привет! На сборку щита нужно:\nLegrand 027005 — две штуки\n027028 — одну штуку\nи ещё 027220, количество по нему уточню после обмера.\nПроверь что есть, если чего нет — сначала покажи замену.\nНужно до конца недели, получится ли — уточним у магазина.\n",
        encoding="utf-8",
    )
    manifest = {
        "version": 2,
        "kind": "realistic_procurement_documents_and_original_catalog_photo",
        "provenance": "Document layouts, project names and quantities are constructed, not customer records. Product codes come from partner catalogue. delivery_note_image.jpg is a raster rendering of the constructed form; product_photo.jpg is unchanged from the published ekt.kz photograph.",
        "requirements": "HackAlem Word sections 3–7, 8 (manager), 9; UC-01…05",
        "prices_and_stock": "Always read from live catalogue; document prices and desired deadlines are not store facts",
        "cart_before_explicit_confirmation": "unchanged",
        "files": [],
    }
    for case in CASES:
        f = OUT / case["path"]
        manifest["files"].append(
            {
                **case,
                "size_bytes": f.stat().st_size,
                "sha256": hashlib.sha256(f.read_bytes()).hexdigest(),
            }
        )
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    print(
        f"Generated {len(CASES)} distinct inputs and per-file expected semantics; no AI answers generated"
    )


if __name__ == "__main__":
    main()
